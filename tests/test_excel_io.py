# tests/test_excel_io.py
"""Ekspor/impor Excel generik (etl/excel_io.py, /api/excel)."""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import text

from etl import excel_io as xio

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx(header, rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(client, entity, content, dry_run=False):
    return client.post(f"/api/excel/{entity}/import" + ("?dry_run=true" if dry_run else ""), content=content,
                       headers={"Content-Type": XLSX})


class TestExport:
    @pytest.mark.parametrize("code", list(xio.ENTITIES))
    def test_every_entity_exports_as_admin(self, make_client, synthetic_aoi, code):
        r = make_client("ADMIN").get(f"/api/excel/{code}.xlsx?date_from=2024-01-01&date_to=2024-12-31")
        assert r.status_code == 200, r.text
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Data", "Petunjuk"]
        assert wb["Data"].max_column >= 2

    def test_role_per_entity(self, make_client, synthetic_aoi, db_client):
        user = make_client("USER")
        assert user.get("/api/excel/alerts.xlsx").status_code == 200
        r = user.get("/api/excel/disasters.xlsx")
        assert r.status_code == 403 and r.json()["code"] == "ENTITY_FORBIDDEN"
        assert make_client("ANALYST").get("/api/excel/datasets.xlsx").status_code == 403
        assert make_client("DATA_ENGINEER").get("/api/excel/audit.xlsx").status_code == 403
        codes = {i["code"] for i in user.get("/api/excel").json()["items"]}
        assert "alerts" in codes and "audit" not in codes and "disasters" not in codes
        with db_client.session() as sess:
            assert sess.scalar(text("SELECT count(*) FROM user_activity_logs WHERE action = 'DOWNLOAD_XLSX'")) >= 1

    def test_kecamatan_export_content(self, make_client, synthetic_aoi):
        wb = load_workbook(io.BytesIO(make_client("USER").get("/api/excel/kecamatan.xlsx").content))
        rows = list(wb["Data"].iter_rows(values_only=True))
        assert rows[0][:3] == ("pcode", "region_name", "in_aoi")
        assert {"TST001", "TST002", "TST003"} <= {r[0] for r in rows[1:]}

    def test_unknown_entity(self, make_client):
        assert make_client("ADMIN").get("/api/excel/nope.xlsx").status_code == 404


class TestImport:
    HEADER = ["tanggal", "jenis", "kecamatan", "desa", "lat", "lon", "keterangan", "sumber", "terverifikasi"]

    def test_template_and_roundtrip_disasters(self, make_client, synthetic_aoi):
        analyst = make_client("ANALYST")
        tpl = load_workbook(io.BytesIO(analyst.get("/api/excel/disasters/template.xlsx").content))
        assert [c.value for c in tpl["Data"][1]][:4] == ["tanggal", "tanggal_selesai", "jenis", "kecamatan"]
        content = _xlsx(self.HEADER, [
            ["2023-03-01", "BANJIR", "TST001", "Desa X", -6.55, 106.1, "Banjir uji impor Excel satu", "GMLS", "ya"],
            ["2023-03-02", "LONGSOR", "Kecamatan 2", None, None, None, "Longsor uji impor Excel dua", "MEDIA", None]])
        dry = _post(analyst, "disasters", content, dry_run=True)
        assert dry.status_code == 200 and dry.json()["inserted"] == 2 and dry.json()["dry_run"] is True
        r = _post(analyst, "disasters", content)
        assert r.status_code == 200 and r.json()["inserted"] == 2, r.text
        assert _post(analyst, "disasters", content).json()["duplicates"] == 2
        # Ekspor -> impor ulang: format kolom sama, semua terdeteksi duplikat.
        exported = analyst.get("/api/excel/disasters.xlsx?date_from=2023-03-01&date_to=2023-03-02").content
        again = _post(analyst, "disasters", exported)
        assert again.status_code == 200 and again.json()["duplicates"] == 2 and again.json()["inserted"] == 0

    def test_invalid_row_rejects_whole_file(self, make_client, synthetic_aoi, db_client):
        content = _xlsx(self.HEADER, [
            ["2023-04-01", "BANJIR", "TST001", None, None, None, "Baris yang benar untuk uji", "GMLS", None],
            ["2023-04-02", "BANJIR", "TIDAKADA", None, None, None, "Kecamatan tidak dikenal", "GMLS", None],
            ["2023-04-03", "BANJIR", "TST002", None, None, None, "pendek", "XYZ", None]])
        r = _post(make_client("ANALYST"), "disasters", content)
        assert r.status_code == 422 and r.json()["code"] == "IMPORT_ROWS_INVALID"
        assert len(r.json()["errors"]) == 2 and r.json()["errors"][0].startswith("line 3")
        with db_client.session() as sess:
            assert sess.scalar(text("SELECT count(*) FROM disaster_events WHERE event_date = '2023-04-01'")) == 0

    def test_bad_file_and_permissions(self, make_client):
        assert _post(make_client("ANALYST"), "disasters", b"not an excel file").json()["code"] == "INVALID_IMPORT_FILE"
        assert _post(make_client("USER"), "disasters", _xlsx(["x"], [[1]])).status_code == 403
        assert _post(make_client("ADMIN"), "alerts", _xlsx(["x"], [[1]])).json()["code"] == "IMPORT_NOT_SUPPORTED"

    def test_alert_rules_and_types_upsert(self, make_client, db_client):
        admin = make_client("ADMIN")
        r = _post(admin, "disaster_types", _xlsx(["type_code", "type_name", "is_active"],
                                                 [["ROB", "Banjir rob", "ya"], ["BANJIR", "Banjir (diperbarui)", None]]))
        assert r.status_code == 200 and (r.json()["inserted"], r.json()["updated"]) == (1, 1), r.text
        hdr = ["rule_code", "disaster_type_code", "band_code", "comparator", "threshold_value", "severity",
               "reference_source", "is_active"]
        r = _post(admin, "alert_rules", _xlsx(hdr, [["XLS_ROB_24H", "ROB", "RAIN_24H", ">=", 80, "WARNING", "uji", "ya"]]))
        assert r.status_code == 200 and r.json()["inserted"] == 1, r.text
        bad = _post(admin, "alert_rules", _xlsx(hdr, [["XLS_BAD", "ROB", "RAIN_24H", ">=", None, "INFO", "uji", "ya"]]))
        assert bad.status_code == 422 and "threshold" in bad.json()["errors"][0]
        assert _post(make_client("ANALYST"), "alert_rules", _xlsx(hdr, [])).status_code == 403
        with db_client.session() as sess:
            assert sess.scalar(text("SELECT type_name FROM disaster_types WHERE type_code = 'BANJIR'")) == "Banjir (diperbarui)"
            sess.execute(text("UPDATE disaster_types SET type_name = 'Banjir' WHERE type_code = 'BANJIR'"))
            sess.execute(text("DELETE FROM alert_rules WHERE rule_code = 'XLS_ROB_24H'"))

    def test_kecamatan_aoi_and_settings(self, make_client, synthetic_aoi, db_client):
        admin = make_client("ADMIN")
        r = _post(admin, "kecamatan", _xlsx(["pcode", "in_aoi"], [["TST003", "tidak"]]))
        assert r.status_code == 200 and r.json()["updated"] == 1
        with db_client.session() as sess:
            assert sess.scalar(text("SELECT in_aoi FROM administrative_regions WHERE pcode = 'TST003'")) is False
        assert _post(admin, "kecamatan", _xlsx(["pcode", "in_aoi"], [["TST003", "ya"]])).status_code == 200
        r = _post(admin, "app_settings", _xlsx(["setting_key", "setting_value"], [["live.retention_default", "8"]]))
        assert r.status_code == 200, r.text
        with db_client.session() as sess:
            assert sess.scalar(text("SELECT setting_value FROM app_settings WHERE setting_key = 'live.retention_default'")) == 8
            sess.execute(text("UPDATE app_settings SET setting_value = '6' WHERE setting_key = 'live.retention_default'"))
        bad = _post(admin, "app_settings", _xlsx(["setting_key", "setting_value"], [["tidak.ada", "1"]]))
        assert bad.status_code == 422
