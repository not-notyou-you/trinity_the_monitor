# etl/excel_io.py
"""Ekspor/impor Excel (.xlsx) generik untuk hampir semua data Monitor (Tahap 3).

Satu registri ``ENTITIES`` menentukan, per jenis data:

* ``export_role``  role minimum untuk mengunduh (sama dengan endpoint baca
  yang setara); kueri berjalan di bawah role pemanggil, jadi GRANT/RLS
  PostgreSQL tetap berlaku.
* ``import_role`` + ``importer``  bila data itu boleh dimasukkan lewat Excel.
  Impor selalu semua-atau-tidak-sama-sekali (pemanggil membungkusnya dalam
  SAVEPOINT) dan melaporkan kesalahan per nomor baris; ``dry_run`` hanya
  memvalidasi.

Berkas berisi sheet ``Data`` (baris 1 = kode kolom) dan sheet ``Petunjuk``
(arti kolom, wajib/tidak, nilai yang diizinkan, filter ekspor). Templat impor
= sheet Data kosong + Petunjuk. Waktu (timestamptz) ditulis dalam WIB tanpa
zona karena Excel tidak menyimpan zona waktu; tanggal hidromet = hari UTC.
"""

from __future__ import annotations

import io
import json
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Callable
from zoneinfo import ZoneInfo

from sqlalchemy import text

WIB = ZoneInfo("Asia/Jakarta")
MAX_IMPORT_ROWS = 20000
MAX_EXPORT_ROWS = 200000
DATA_SHEET = "Data"
HELP_SHEET = "Petunjuk"


@dataclass(frozen=True)
class Col:
    key: str
    help: str = ""
    required: bool = False
    allowed: str = ""


@dataclass
class Entity:
    code: str
    title: str
    export_role: str
    sql: str                       # SELECT ...; filter tanggal ditambahkan lewat date_column
    date_column: str | None = None
    order_by: str = "1"
    export_help: dict[str, str] = field(default_factory=dict)
    import_role: str | None = None
    import_columns: list[Col] = field(default_factory=list)
    importer: Callable | None = None   # (sess, header, rows, user_id, dry_run) -> summary


# --- importer -----------------------------------------------------------------------

def _summary(inserted=0, updated=0, duplicates=0, errors=None) -> dict:
    return {"inserted": inserted, "updated": updated, "duplicates": duplicates, "errors": errors or []}


def _check_header(header: list[str], cols: list[Col]) -> None:
    have = {h for h in header if h}
    missing = [c.key for c in cols if c.required and c.key not in have]
    unknown = sorted(have - {c.key for c in cols})
    if missing:
        raise ValueError("missing required column(s): " + ", ".join(missing))
    if unknown:
        raise ValueError("unknown column(s): " + ", ".join(unknown))


def _bool(v) -> bool | None:
    s = str(v or "").strip().lower()
    if s == "":
        return None
    if s in ("1", "true", "ya", "y", "yes"):
        return True
    if s in ("0", "false", "tidak", "t", "n", "no"):
        return False
    raise ValueError(f"not a yes/no value: {v!r}")


def _num(v) -> float | None:
    s = str(v or "").strip().replace(",", ".")
    if s == "":
        return None
    return float(s)


def import_disasters(sess, header, rows, user_id, dry_run):
    from etl.disasters import import_rows
    return import_rows(sess, header, rows, user_id, dry_run)


def import_disaster_types(sess, header, rows, user_id, dry_run):
    _check_header(header, ENTITIES["disaster_types"].import_columns)
    ins = upd = 0
    errors = []
    for n, r in rows:
        try:
            code = str(r.get("type_code") or "").strip().upper()
            name = str(r.get("type_name") or "").strip()
            if not code or not name:
                raise ValueError("type_code and type_name are required")
            active = _bool(r.get("is_active"))
            existed = sess.scalar(text("SELECT 1 FROM disaster_types WHERE type_code = :c"), {"c": code})
            sess.execute(text("""
                INSERT INTO disaster_types (type_code, type_name, category, indicator_bands, is_active)
                VALUES (:c, :n, COALESCE(NULLIF(:cat, ''), 'HIDROMETEOROLOGI'), NULLIF(:ib, ''), COALESCE(:a, true))
                ON CONFLICT (type_code) DO UPDATE SET type_name = EXCLUDED.type_name,
                    category = EXCLUDED.category, indicator_bands = EXCLUDED.indicator_bands,
                    is_active = COALESCE(:a, disaster_types.is_active)"""),
                {"c": code, "n": name, "cat": str(r.get("category") or ""), "ib": str(r.get("indicator_bands") or ""),
                 "a": active})
            upd += bool(existed)
            ins += not existed
        except Exception as exc:
            errors.append(f"line {n}: {exc}")
    return _summary(ins, upd, 0, errors)


def import_alert_rules(sess, header, rows, user_id, dry_run):
    _check_header(header, ENTITIES["alert_rules"].import_columns)
    ins = upd = 0
    errors = []
    for n, r in rows:
        try:
            code = str(r.get("rule_code") or "").strip().upper()
            if not code:
                raise ValueError("rule_code is required")
            cmp_ = str(r.get("comparator") or ">=").strip()
            sev = str(r.get("severity") or "").strip().upper()
            if cmp_ not in (">=", ">", "<=", "<"):
                raise ValueError("comparator must be >=, >, <= or <")
            if sev not in ("INFO", "WARNING", "CRITICAL"):
                raise ValueError("severity must be INFO, WARNING or CRITICAL")
            thr = _num(r.get("threshold_value"))
            active = _bool(r.get("is_active"))
            active = True if active is None else active
            if active and thr is None:
                raise ValueError("an active rule needs threshold_value")
            ids = sess.execute(text("""SELECT (SELECT disaster_type_id FROM disaster_types WHERE type_code = :t) AS dt,
                                              (SELECT band_id FROM spectral_bands WHERE band_code = :b) AS band"""),
                               {"t": str(r.get("disaster_type_code") or "").strip().upper(),
                                "b": str(r.get("band_code") or "").strip().upper()}).one()
            if ids.dt is None or ids.band is None:
                raise ValueError("unknown disaster_type_code or band_code")
            ref = str(r.get("reference_source") or "").strip()
            if len(ref) < 2:
                raise ValueError("reference_source is required")
            existed = sess.scalar(text("SELECT 1 FROM alert_rules WHERE rule_code = :c"), {"c": code})
            sess.execute(text("""
                INSERT INTO alert_rules (rule_code, disaster_type_id, band_id, comparator, threshold_value, severity,
                                         reference_source, is_active, updated_by)
                VALUES (:c, :dt, :b, :cmp, :thr, :sev, :ref, :a, :u)
                ON CONFLICT (rule_code) DO UPDATE SET disaster_type_id = EXCLUDED.disaster_type_id,
                    band_id = EXCLUDED.band_id, comparator = EXCLUDED.comparator,
                    threshold_value = EXCLUDED.threshold_value, severity = EXCLUDED.severity,
                    reference_source = EXCLUDED.reference_source, is_active = EXCLUDED.is_active,
                    updated_by = EXCLUDED.updated_by, updated_at = now()"""),
                {"c": code, "dt": ids.dt, "b": ids.band, "cmp": cmp_, "thr": thr, "sev": sev, "ref": ref,
                 "a": active, "u": user_id})
            upd += bool(existed)
            ins += not existed
        except Exception as exc:
            errors.append(f"line {n}: {exc}")
    return _summary(ins, upd, 0, errors)


def import_kecamatan_aoi(sess, header, rows, user_id, dry_run):
    """Kolom in_aoi per pcode; kecamatan yang tidak disebut tidak berubah.
    ROI AOI + dataset HYDROMET_AOI dibangun ulang."""
    from etl import regions as rg
    _check_header(header, ENTITIES["kecamatan"].import_columns)
    upd = 0
    errors = []
    for n, r in rows:
        try:
            pcode = str(r.get("pcode") or "").strip().upper()
            flag = _bool(r.get("in_aoi"))
            if not pcode or flag is None:
                raise ValueError("pcode and in_aoi are required")
            hit = sess.execute(text("""UPDATE administrative_regions SET in_aoi = :f
                                       WHERE upper(pcode) = :p AND admin_level = 3 RETURNING region_id"""),
                               {"f": flag, "p": pcode}).first()
            if hit is None:
                raise ValueError(f"kecamatan {pcode} not found")
            upd += 1
        except Exception as exc:
            errors.append(f"line {n}: {exc}")
    if not errors:
        if not sess.scalar(text("SELECT count(*) FROM administrative_regions WHERE in_aoi")):
            errors.append("the AOI must keep at least one kecamatan")
        else:
            rg.rebuild_monitor_aoi(sess)
    return _summary(0, upd, 0, errors)


def import_app_settings(sess, header, rows, user_id, dry_run):
    _check_header(header, ENTITIES["app_settings"].import_columns)
    upd = 0
    errors = []
    for n, r in rows:
        try:
            key = str(r.get("setting_key") or "").strip()
            raw = r.get("setting_value")
            raw = "" if raw is None else str(raw).strip()
            try:
                value = json.loads(raw)
            except json.JSONDecodeError:
                value = raw          # teks biasa -> string JSON
            hit = sess.execute(text("""UPDATE app_settings SET setting_value = CAST(:v AS jsonb), updated_by = :u,
                                              updated_at = now() WHERE setting_key = :k RETURNING setting_key"""),
                               {"v": json.dumps(value), "u": user_id, "k": key}).first()
            if hit is None:
                raise ValueError(f"unknown setting_key {key!r} (new keys are added by the seed, not by import)")
            upd += 1
        except Exception as exc:
            errors.append(f"line {n}: {exc}")
    return _summary(0, upd, 0, errors)


# --- registri -----------------------------------------------------------------------

def _disaster_cols() -> list[Col]:
    from etl.disasters import INFO_SOURCES
    return [
        Col("tanggal", "Tanggal kejadian (YYYY-MM-DD).", True),
        Col("tanggal_selesai", "Tanggal selesai, kosong bila satu hari."),
        Col("jenis", "Kode jenis bencana (lihat ekspor disaster_types).", True, "BANJIR, BANJIR_BANDANG, LONGSOR, KEKERINGAN, …"),
        Col("kecamatan", "Nama atau P-code kecamatan COD-AB.", True),
        Col("desa", "Nama desa/kelurahan."),
        Col("lat", "Lintang (desimal, WGS84)."), Col("lon", "Bujur (desimal, WGS84)."),
        Col("keterangan", "Uraian kejadian (10–4000 karakter, tanpa data pribadi).", True),
        Col("dampak", "Ringkasan dampak."),
        Col("sumber", "Sumber informasi.", True, ", ".join(INFO_SOURCES)),
        Col("referensi", "Tautan/nomor dokumen sumber."),
        Col("terverifikasi", "ya/tidak.", False, "ya, tidak"),
    ]


ENTITIES: dict[str, Entity] = {e.code: e for e in [
    Entity("kecamatan", "Kecamatan Kabupaten Lebak (COD-AB) dan status AOI", "USER",
           """SELECT r.pcode, r.region_name, r.in_aoi, r.area_km2, p.pcode AS kabupaten_pcode, r.source_dataset
              FROM administrative_regions r LEFT JOIN administrative_regions p ON p.region_id = r.parent_region_id
              WHERE r.admin_level = 3""", order_by="r.region_name",
           import_role="ADMIN", importer=import_kecamatan_aoi,
           import_columns=[Col("pcode", "P-code COD-AB kecamatan.", True),
                           Col("in_aoi", "ya = termasuk AOI GMLS.", True, "ya, tidak"),
                           Col("region_name", "Diabaikan saat impor (informasi)."),
                           Col("area_km2", "Diabaikan saat impor."), Col("kabupaten_pcode", "Diabaikan saat impor."),
                           Col("source_dataset", "Diabaikan saat impor.")]),
    Entity("observations", "Observasi harian per kecamatan (region_observations)", "ANALYST",
           """SELECT o.obs_date AS tanggal_utc, r.pcode, r.region_name, b.band_code, b.unit, o.value,
                     o.valid_fraction, o.run_type, o.source_product_id, o.computed_at
              FROM region_observations o JOIN spectral_bands b ON b.band_id = o.band_id
              JOIN administrative_regions r ON r.region_id = o.region_id""",
           date_column="o.obs_date", order_by="o.obs_date, r.region_name, b.band_code"),
    Entity("hujan_harian", "Hujan harian per kecamatan + kategori BMKG (v_hujan_harian_kecamatan)", "ANALYST",
           "SELECT * FROM v_hujan_harian_kecamatan v", date_column="v.obs_date", order_by="v.obs_date, v.region_name"),
    Entity("alerts", "Alert hujan (alert_events)", "USER",
           """SELECT a.alert_id, a.observation_date, r.pcode, r.region_name, ru.rule_code, a.severity,
                     a.observed_value, a.threshold_value, a.triggered_at, a.acknowledged_at, a.ack_note
              FROM alert_events a JOIN administrative_regions r ON r.region_id = a.region_id
              JOIN alert_rules ru ON ru.rule_id = a.rule_id""",
           date_column="a.observation_date", order_by="a.observation_date, r.region_name"),
    Entity("alert_rules", "Aturan alert", "USER",
           """SELECT ru.rule_code, dt.type_code AS disaster_type_code, b.band_code, ru.comparator, ru.threshold_value,
                     ru.severity, ru.reference_source, ru.is_active
              FROM alert_rules ru JOIN disaster_types dt ON dt.disaster_type_id = ru.disaster_type_id
              JOIN spectral_bands b ON b.band_id = ru.band_id""", order_by="ru.rule_code",
           import_role="ADMIN", importer=import_alert_rules,
           import_columns=[Col("rule_code", "Kode unik (huruf besar, angka, _). Baris dengan kode yang ada diperbarui.", True),
                           Col("disaster_type_code", "Kode jenis bencana.", True),
                           Col("band_code", "Band observasi.", True, "RAIN_24H, RAIN_72H, RAIN_7D, RAIN_30D, FLOOD, NDVI, NDWI"),
                           Col("comparator", "Pembanding.", False, ">=, >, <=, <"),
                           Col("threshold_value", "Ambang (wajib bila aktif)."),
                           Col("severity", "Tingkat.", True, "INFO, WARNING, CRITICAL"),
                           Col("reference_source", "Rujukan ambang.", True),
                           Col("is_active", "ya/tidak (default ya).", False, "ya, tidak")]),
    Entity("disaster_types", "Master jenis bencana", "USER",
           "SELECT type_code, type_name, category, indicator_bands, is_active FROM disaster_types",
           order_by="type_code", import_role="ADMIN", importer=import_disaster_types,
           import_columns=[Col("type_code", "Kode unik. Baris dengan kode yang ada diperbarui.", True),
                           Col("type_name", "Nama tampilan.", True), Col("category", "Default HIDROMETEOROLOGI."),
                           Col("indicator_bands", "Teks informatif band indikator."),
                           Col("is_active", "ya/tidak.", False, "ya, tidak")]),
    Entity("disasters", "Kejadian bencana", "ANALYST",
           """SELECT e.event_date AS tanggal, e.event_end_date AS tanggal_selesai, dt.type_code AS jenis,
                     r.pcode AS kecamatan, e.village_name AS desa, ST_Y(e.location) AS lat, ST_X(e.location) AS lon,
                     e.description AS keterangan, e.impact_summary AS dampak, e.info_source AS sumber,
                     e.source_reference AS referensi, CASE WHEN e.is_verified THEN 'ya' ELSE 'tidak' END AS terverifikasi
              FROM disaster_events e JOIN disaster_types dt ON dt.disaster_type_id = e.disaster_type_id
              JOIN administrative_regions r ON r.region_id = e.region_id
              WHERE e.deleted_at IS NULL""",
           date_column="e.event_date", order_by="e.event_date, r.pcode",
           import_role="ANALYST", importer=import_disasters, import_columns=_disaster_cols()),
    Entity("disaster_rain", "Kejadian bencana + hujan H-0..H-2 (v_kejadian_dan_hujan)", "ANALYST",
           "SELECT * FROM v_kejadian_dan_hujan v", date_column="v.event_date", order_by="v.event_date"),
    Entity("alert_evaluation", "Evaluasi alert: hit / miss / false alarm (v_evaluasi_alert)", "ANALYST",
           "SELECT * FROM v_evaluasi_alert v", date_column="v.ref_date", order_by="v.ref_date"),
    Entity("live_scenes", "Scene Live Sentinel-1 + metrik (live_scenes, live_scene_metrics)", "USER",
           """SELECT a.name AS area, s.scene_date, s.status, b.band_code, m.metric_name, m.value, m.source_date
              FROM live_scenes s JOIN live_areas a ON a.area_id = s.area_id
              LEFT JOIN live_scene_metrics m ON m.live_scene_id = s.live_scene_id
              LEFT JOIN spectral_bands b ON b.band_id = m.band_id""",
           date_column="s.scene_date", order_by="a.name, s.scene_date, b.band_code, m.metric_name"),
    Entity("datasets", "Dataset Katalog", "DATA_ENGINEER",
           """SELECT dataset_id, name, location_label, date_start, date_end, fusion_strategy, status, total_scenes,
                     completed_scenes, failed_scenes, total_size_bytes, created_by, created_at
              FROM datasets d WHERE NOT d.is_system AND d.deleted_at IS NULL""",
           date_column="d.date_start", order_by="d.dataset_id"),
    Entity("s1_scenes", "Scene Sentinel-1 (satellite_scenes)", "DATA_ENGINEER",
           """SELECT scene_id, product_identifier, acquisition_datetime, orbit_number, relative_orbit, orbit_direction,
                     is_valid, invalid_reason FROM satellite_scenes s""",
           date_column="(s.acquisition_datetime AT TIME ZONE 'UTC')::date", order_by="s.acquisition_datetime"),
    Entity("nasa_scenes", "Granule MODIS/GPM (nasa_scenes)", "DATA_ENGINEER",
           """SELECT nasa_scene_id, source, product_short_name, tile_id, acquisition_date, run_type, is_valid,
                     invalid_reason FROM nasa_scenes n""",
           date_column="n.acquisition_date", order_by="n.acquisition_date, n.source"),
    Entity("products", "Produk data (data_products)", "DATA_ENGINEER",
           """SELECT product_id, dataset_id, source, product_tier::text AS product_tier, band_name, file_name,
                     file_size_mb, data_hash_sha256, is_valid, is_latest, created_at FROM data_products p""",
           date_column="(p.created_at AT TIME ZONE 'Asia/Jakarta')::date", order_by="p.product_id"),
    Entity("quality_summary", "Ringkasan kualitas per sumber per minggu (v_ringkasan_kualitas)", "DATA_ENGINEER",
           "SELECT * FROM v_ringkasan_kualitas v", date_column="v.week_start", order_by="v.week_start, v.source_code"),
    Entity("completeness", "Kelengkapan data per sumber (v_kelengkapan_data)", "DATA_ENGINEER",
           "SELECT * FROM v_kelengkapan_data v", date_column="v.data_date", order_by="v.source_code, v.data_date"),
    Entity("downloads", "Log unduhan (v_log_unduhan)", "ADMIN",
           "SELECT * FROM v_log_unduhan v", date_column="(v.logged_at AT TIME ZONE 'Asia/Jakarta')::date",
           order_by="v.logged_at"),
    Entity("audit", "Audit perubahan data (audit_log)", "ADMIN",
           """SELECT audit_id, changed_at, table_name, row_pk, operation, app_user_id, db_user, changed_columns,
                     old_data, new_data FROM audit_log a""",
           date_column="(a.changed_at AT TIME ZONE 'Asia/Jakarta')::date", order_by="a.audit_id"),
    Entity("app_settings", "Pengaturan aplikasi (app_settings)", "ADMIN",
           "SELECT setting_key, setting_value, description, updated_at FROM app_settings", order_by="setting_key",
           import_role="ADMIN", importer=import_app_settings,
           import_columns=[Col("setting_key", "Kunci yang sudah ada.", True),
                           Col("setting_value", "Nilai JSON (angka, \"teks\", true/false).", True),
                           Col("description", "Diabaikan saat impor."), Col("updated_at", "Diabaikan saat impor.")]),
]}


# --- ekspor -------------------------------------------------------------------------

def _cell(v):
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, datetime):
        if v.tzinfo is not None:
            v = v.astimezone(WIB).replace(tzinfo=None)
        return v
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False, default=str)
    if isinstance(v, (bytes, memoryview)):
        return None
    if v is not None and not isinstance(v, (str, int, float, bool, date)):
        return str(v)
    return v


def query_rows(sess, entity: Entity, date_from: date | None, date_to: date | None) -> tuple[list[str], list]:
    sql, params = entity.sql, {}
    conds = []
    if entity.date_column and date_from:
        conds.append(f"{entity.date_column} >= :a")
        params["a"] = date_from
    if entity.date_column and date_to:
        conds.append(f"{entity.date_column} <= :b")
        params["b"] = date_to
    if conds:
        # SQL registri tidak memakai subquery ber-WHERE, jadi cukup cek WHERE di level atas.
        joiner = " AND " if " where " in " ".join(sql.lower().split()) else " WHERE "
        sql = sql + joiner + " AND ".join(conds)
    res = sess.execute(text(f"{sql} ORDER BY {entity.order_by} LIMIT {MAX_EXPORT_ROWS + 1}"), params)
    return list(res.keys()), res.all()


def _help_sheet(wb, entity: Entity, columns: list[str], note: list[tuple[str, str]], for_import: bool) -> None:
    from openpyxl.styles import Font
    ws = wb.create_sheet(HELP_SHEET)
    ws.append([entity.title])
    ws["A1"].font = Font(bold=True, size=13)
    for k, v in note:
        ws.append([k, v])
    ws.append([])
    if entity.importer and entity.import_columns:
        ws.append(["Impor", f"didukung (role minimum {entity.import_role}); semua baris atau tidak sama sekali"])
        ws.append([])
        ws.append(["Kolom", "Wajib", "Keterangan", "Nilai yang diizinkan"])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        for c in entity.import_columns:
            ws.append([c.key, "ya" if c.required else "", c.help, c.allowed])
    else:
        ws.append(["Impor", "tidak didukung untuk data ini (hanya ekspor)"])
        ws.append([])
        ws.append(["Kolom"])
        for col in columns:
            ws.append([col, entity.export_help.get(col, "")])
    ws.append([])
    ws.append(["Waktu (timestamp) dalam WIB; tanggal hidromet (obs_date) = hari UTC. Sel kosong = tidak ada data."])
    for col, width in (("A", 26), ("B", 12), ("C", 70), ("D", 50)):
        ws.column_dimensions[col].width = width


def _data_sheet(wb, columns: list[str], rows) -> None:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws = wb.active
    ws.title = DATA_SHEET
    ws.append(columns)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4C72B0")
    for r in rows:
        ws.append([_cell(v) for v in r])
    ws.freeze_panes = "A2"
    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(45, len(col) + 4))


def export_xlsx(sess, code: str, date_from: date | None = None, date_to: date | None = None) -> tuple[bytes, int]:
    """(isi .xlsx, jumlah baris)."""
    from openpyxl import Workbook
    entity = ENTITIES[code]
    columns, rows = query_rows(sess, entity, date_from, date_to)
    truncated = len(rows) > MAX_EXPORT_ROWS
    rows = rows[:MAX_EXPORT_ROWS]
    wb = Workbook()
    _data_sheet(wb, columns, rows)
    note = [("Diekspor", datetime.now(WIB).strftime("%Y-%m-%d %H:%M WIB")), ("Jumlah baris", str(len(rows)))]
    if entity.date_column:
        note.append(("Filter tanggal", f"{date_from or '—'} s.d. {date_to or '—'}"))
    if truncated:
        note.append(("Peringatan", f"dipotong pada {MAX_EXPORT_ROWS} baris; persempit rentang tanggal"))
    _help_sheet(wb, entity, columns, note, False)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(rows)


def template_xlsx(code: str) -> bytes:
    from openpyxl import Workbook
    entity = ENTITIES[code]
    if not entity.importer:
        raise ValueError(f"{code} cannot be imported")
    wb = Workbook()
    _data_sheet(wb, [c.key for c in entity.import_columns], [])
    _help_sheet(wb, entity, [], [("Templat impor", "isi sheet Data mulai baris 2; jangan ubah baris 1")], True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- impor --------------------------------------------------------------------------

def _plain(v):
    """Nilai sel -> teks seperti CSV (tanggal ISO, angka tanpa .0 berlebih)."""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, bool):
        return "ya" if v else "tidak"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_xlsx(content: bytes) -> tuple[list[str], list[tuple[int, dict]]]:
    """(header, [(nomor baris Excel, {kolom: teks})]) dari sheet Data (atau sheet pertama)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"not a valid .xlsx file ({type(exc).__name__})")
    ws = wb[DATA_SHEET] if DATA_SHEET in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    try:
        header = [_plain(h).lower() for h in next(it)]
    except StopIteration:
        raise ValueError("the sheet is empty")
    rows = []
    for n, values in enumerate(it, start=2):
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        rows.append((n, {h: _plain(v) for h, v in zip(header, values) if h}))
        if len(rows) > MAX_IMPORT_ROWS:
            raise ValueError(f"at most {MAX_IMPORT_ROWS} rows per import")
    wb.close()
    return [h for h in header if h], rows


def import_xlsx(sess, code: str, content: bytes, user_id: int | None, dry_run: bool = False) -> dict:
    """Validasi + tulis dalam SAVEPOINT; ada error atau dry_run -> dibatalkan."""
    entity = ENTITIES[code]
    if not entity.importer:
        raise ValueError(f"{code} cannot be imported")
    header, rows = read_xlsx(content)
    if not rows:
        return _summary()
    nested = sess.begin_nested()
    try:
        summary = entity.importer(sess, header, rows, user_id, dry_run)
    except Exception:
        nested.rollback()
        raise
    if summary["errors"] or dry_run:
        nested.rollback()
        if summary["errors"]:
            summary.update(inserted=0, updated=0)
    else:
        nested.commit()
    summary["rows"] = len(rows)
    summary["dry_run"] = dry_run
    return summary
